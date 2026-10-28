#!/usr/bin/env python3
"""Build the downloadable Excel templates into public/downloads/.

    python scripts/templates_xlsx.py

Then recalculate each file so cached values exist (see README):
    python <xlsx skill>/scripts/recalc.py public/downloads/<file>.xlsx
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parent.parent / "public" / "downloads"
F = "Arial"
VIOLET, NAVY, CYAN = "5B3FD6", "1A1640", "0E9AA7"
H1 = Font(name=F, size=16, bold=True, color=NAVY)
H2 = Font(name=F, size=11, bold=True, color="FFFFFF")
B = Font(name=F, size=10)
BB = Font(name=F, size=10, bold=True)
DIM = Font(name=F, size=9, italic=True, color="666666")
INPUT = Font(name=F, size=10, color="0000FF")
LINK = Font(name=F, size=10, color="008000")
HEAD_FILL = PatternFill("solid", fgColor=VIOLET)
SUB_FILL = PatternFill("solid", fgColor="E9E5FB")
IN_FILL = PatternFill("solid", fgColor="FFF6BF")
TOT_FILL = PatternFill("solid", fgColor="DDF5F7")
THIN = Side(style="thin", color="CCCCCC")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
NUM = '#,##0;(#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'
FOOT = "Shopfloor Matters · shopfloormatters.com · Free to use and adapt. General information, not professional advice."


def title(ws, text, sub=None, width_cols=6):
    ws["A1"] = text
    ws["A1"].font = H1
    if sub:
        ws["A2"] = sub
        ws["A2"].font = DIM
        ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width_cols)
        ws.row_dimensions[2].height = 42
    ws.sheet_view.showGridLines = False


def header(ws, row, labels, col=1):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=lab)
        c.font, c.fill, c.border = H2, HEAD_FILL, BOX
        c.alignment = Alignment(wrap_text=True, vertical="center")


def inp(c, value, fmt=None):
    c.value = value
    c.font, c.fill, c.border = INPUT, IN_FILL, BOX
    if fmt:
        c.number_format = fmt


def cell(ws, ref, value, font=B, fmt=None, fill=None):
    c = ws[ref]
    c.value, c.font, c.border = value, font, BOX
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    return c


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def legend(ws, row):
    ws.cell(row=row, column=1, value="Key").font = BB
    c = ws.cell(row=row + 1, column=1, value="Yellow cell, blue text")
    c.font, c.fill = INPUT, IN_FILL
    ws.cell(row=row + 1, column=2, value="Your input. Replace the example values with your own.").font = B
    ws.cell(row=row + 2, column=1, value="Black text").font = B
    ws.cell(row=row + 2, column=2, value="Calculated. Leave these alone.").font = B
    ws.cell(row=row + 3, column=1, value="Green text").font = LINK
    ws.cell(row=row + 3, column=2, value="Pulled from another sheet.").font = B


def start_sheet(wb, name, heading, lines):
    ws = wb.active
    ws.title = name
    title(ws, heading)
    r = 3
    for ln in lines:
        c = ws.cell(row=r, column=1, value=ln)
        c.font = BB if ln.endswith(":") else B
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
        ws.row_dimensions[r].height = 30 if len(ln) > 95 else 16
        r += 1
    legend(ws, r + 1)
    ws.cell(row=r + 6, column=1, value=FOOT).font = DIM
    widths(ws, [24, 14, 14, 14, 14, 14, 14, 14])
    return ws


# ---------------------------------------------------------------- business case
def business_case():
    wb = Workbook()
    start_sheet(wb, "Start here", "MES business case template", [
        "Business cases get approved faster when they show four things clearly:",
        "1. What the problems cost today, in money (Baseline).",
        "2. What will improve, by how much and when (Benefits).",
        "3. Realistic costs and timeline, including your own people, not just licences and services (Costs).",
        "4. How the risk is managed: a phased rollout plan with measures and milestones (Rollout plan).",
        "How to use it:",
        "Fill in Baseline with numbers finance will recognise. Set deliberately modest improvements on Benefits. Fill in every cost line, then read Cash flow.",
        "Make the decision once, for the whole programme, then phase the delivery: by function, by site, or both.",
        "The example values describe a fictional mid-size site so you can see how it works. They are not a forecast. Replace them all.",
        "Every amount is in your own currency. Percentages are entered as percentages (35% not 0.35 is fine, Excel handles both).",
    ])

    ws = wb.create_sheet("Baseline")
    title(ws, "Baseline: what the problems cost today", "Your site as it is now, for the lines or area in scope. Ask finance for revenue, margin and material cost. Measure the rest on the floor for a week if you can.")
    header(ws, 4, ["Input", "Value", "Unit", "Where to get it"])
    rows = [
        ("Annual revenue from the lines in scope", 25000000, NUM, "currency / year", "Finance. Sales value of what these lines make."),
        ("Contribution margin", 0.35, PCT, "% of revenue", "Finance. Revenue minus variable costs."),
        ("Current OEE", 0.55, PCT, "%", "Measured, not estimated. Manual OEE usually flatters you."),
        ("Share of extra output you could sell", 0.5, PCT, "%", "Sales. 100% if demand exceeds capacity, 0% if not."),
        ("Annual material cost on these lines", 11000000, NUM, "currency / year", "Finance or ERP."),
        ("Scrap and rework rate", 0.04, PCT, "% of material", "Quality. Scrap value ÷ material consumed."),
        ("Hours a year spent on paper and re-keying", 9000, NUM, "hours / year", "Time it: forms, typing up, meeting prep, finding records."),
        ("Loaded hourly rate", 32, NUM, "currency / hour", "HR or finance. Wage plus on-costs."),
        ("Average inventory value (raw, WIP, finished)", 4000000, NUM, "currency", "Finance."),
        ("Annual carrying cost of inventory", 0.2, PCT, "% of value", "Finance. Capital, storage, obsolescence. 15% to 25% is common."),
        ("Annual maintenance spend on these lines", 1200000, NUM, "currency / year", "Maintenance or finance. Labour, contractors and spares."),
    ]
    for i, (lab, v, fmt, unit, src) in enumerate(rows, start=5):
        cell(ws, f"A{i}", lab)
        inp(ws[f"B{i}"], v, fmt)
        cell(ws, f"C{i}", unit, DIM)
        cell(ws, f"D{i}", src, DIM).alignment = WRAP
    r = 5 + len(rows) + 1
    cell(ws, f"A{r}", "Cost of scrap and rework today", BB)
    cell(ws, f"B{r}", "=B9*B10", BB, NUM, TOT_FILL)
    cell(ws, f"A{r+1}", "Cost of paper and re-keying today", BB)
    cell(ws, f"B{r+1}", "=B11*B12", BB, NUM, TOT_FILL)
    cell(ws, f"A{r+2}", "Cost of carrying inventory today", BB)
    cell(ws, f"B{r+2}", "=B13*B14", BB, NUM, TOT_FILL)
    cell(ws, f"A{r+3}", "Value of one OEE point (the 1% question)", BB)
    cell(ws, f"B{r+3}", "=IF(B7>0,B5*(0.01/B7)*B8*B6,0)", BB, NUM, TOT_FILL)
    ws[f"B{r+3}"].comment = Comment("Revenue × (1 point ÷ current OEE) × share you can sell × contribution margin.", "Shopfloor Matters")
    widths(ws, [44, 16, 16, 58])

    ws = wb.create_sheet("Benefits")
    title(ws, "Benefits: what improves, and when", "Set improvements at the modest end. A board trusts a case that under-promises. The ramp sets how much of the full benefit you bank each year while the system beds in.")
    header(ws, 4, ["Lever", "Improvement", "How it's measured", "Full-year benefit", "Note"])
    levers = [
        ("Higher OEE (points gained)", 5, '0.0" pts"', "OEE points", "=Baseline!B20*B5", "Uses the value of one OEE point from Baseline."),
        ("Less scrap and rework", 0.1, PCT, "% cut in today's scrap cost", "=Baseline!B17*B6", "A 10% cut in scrap is a modest target."),
        ("Less paper and re-keying", 0.5, PCT, "% of today's paper hours removed", "=Baseline!B18*B7", "Count only hours you'd really redeploy or stop paying for."),
        ("Lower inventory", 0.1, PCT, "% cut in average inventory", "=Baseline!B19*B8", "Carrying cost saved. The one-off cash release is not counted."),
        ("Less reactive maintenance", 0.05, PCT, "% cut in maintenance spend", "=Baseline!B15*B9", "From planned work replacing breakdowns."),
    ]
    for i, (lab, v, fmt, how, f, note) in enumerate(levers, start=5):
        cell(ws, f"A{i}", lab)
        inp(ws[f"B{i}"], v, fmt)
        cell(ws, f"C{i}", how, DIM)
        cell(ws, f"D{i}", f, LINK, NUM)
        cell(ws, f"E{i}", note, DIM).alignment = WRAP
    cell(ws, "A10", "Total full-year benefit", BB)
    cell(ws, "D10", "=SUM(D5:D9)", BB, NUM, TOT_FILL)
    cell(ws, "A12", "Benefit ramp", BB)
    for j, (y, v) in enumerate([("Year 1", 0.4), ("Year 2", 0.8), ("Year 3", 1.0)]):
        cell(ws, f"{get_column_letter(2+j)}11", y, BB)
        inp(ws[f"{get_column_letter(2+j)}12"], v, PCT)
    cell(ws, "A13", "Not counted here", DIM)
    cell(ws, "B13", "Quality escapes avoided, audit time, customer retention, less firefighting. Mention them, but keep them out of the numbers.", DIM).alignment = WRAP
    ws.merge_cells("B13:E13")
    ws.row_dimensions[13].height = 30
    widths(ws, [30, 14, 32, 18, 50])

    ws = wb.create_sheet("Costs")
    title(ws, "Costs: every line, including your own people", "The internal team is the cost people forget. Name the people and the share of their time. Put contingency in on purpose, not as padding.")
    header(ws, 4, ["Cost line", "Year 0 (setup)", "Year 1", "Year 2", "Year 3", "Note"])
    costs = [
        ("Software subscription or licences", 0, 90000, 90000, 90000, "Subscription each year, or licences up front plus maintenance."),
        ("Implementation services", 250000, 30000, 0, 0, "Partner or vendor services. Ask for the plan behind the number."),
        ("Hardware and devices", 80000, 10000, 10000, 10000, "Tablets, terminals, scanners, printers, spares."),
        ("Connectivity and infrastructure", 60000, 0, 0, 0, "Machine connectivity, network, Wi-Fi, edge."),
        ("Internal team", 150000, 80000, 40000, 40000, "People × share of time × loaded cost. Project lead, super users, IT."),
        ("Training", 30000, 10000, 5000, 5000, "Operators on every shift, super users, administrators."),
        ("Running and support", 0, 25000, 25000, 25000, "Hosting, internal support, interface monitoring."),
    ]
    for i, (lab, *vals, note) in enumerate(costs, start=5):
        cell(ws, f"A{i}", lab)
        for j, v in enumerate(vals):
            inp(ws.cell(row=i, column=2 + j), v, NUM)
        cell(ws, f"F{i}", note, DIM).alignment = WRAP
    cell(ws, "A12", "Contingency", BB)
    inp(ws["B13"], 0.15, PCT)
    cell(ws, "A13", "Contingency rate (applied to every year)")
    for j in range(4):
        col = get_column_letter(2 + j)
        cell(ws, f"{col}12", f"=SUM({col}5:{col}11)*$B$13", B, NUM)
        cell(ws, f"{col}14", f"=SUM({col}5:{col}12)", BB, NUM, TOT_FILL)
    cell(ws, "A14", "Total cost", BB)
    widths(ws, [38, 15, 13, 13, 13, 52])

    ws = wb.create_sheet("Cash flow")
    title(ws, "Cash flow, payback and ROI", "Three years, using the ramp on Benefits. Payback is the first year cumulative cash turns positive.")
    header(ws, 4, ["", "Year 0", "Year 1", "Year 2", "Year 3"])
    cell(ws, "A5", "Benefit")
    cell(ws, "B5", 0, B, NUM)
    for j in range(3):
        col = get_column_letter(3 + j)
        cell(ws, f"{col}5", f"=Benefits!$D$10*Benefits!{get_column_letter(2+j)}12", LINK, NUM)
    cell(ws, "A6", "Cost")
    cell(ws, "A7", "Net cash flow", BB)
    cell(ws, "A8", "Cumulative", BB)
    for j in range(4):
        col = get_column_letter(2 + j)
        cell(ws, f"{col}6", f"=-Costs!{col}14", LINK, NUM)
        cell(ws, f"{col}7", f"={col}5+{col}6", BB, NUM)
        cell(ws, f"{col}8", "=B7" if j == 0 else f"={get_column_letter(1+j)}8+{col}7", BB, NUM, TOT_FILL)
    cell(ws, "A10", "Three-year benefit", BB)
    cell(ws, "B10", "=SUM(C5:E5)", BB, NUM)
    cell(ws, "A11", "Three-year cost", BB)
    cell(ws, "B11", "=-SUM(B6:E6)", BB, NUM)
    cell(ws, "A12", "Three-year net benefit", BB)
    cell(ws, "B12", "=B10-B11", BB, NUM, TOT_FILL)
    cell(ws, "A13", "Three-year ROI", BB)
    cell(ws, "B13", "=IF(B11>0,B12/B11,0)", BB, PCT, TOT_FILL)
    cell(ws, "A14", "Payback", BB)
    cell(ws, "B14", '=IF(C8>=0,"Year 1",IF(D8>=0,"Year 2",IF(E8>=0,"Year 3","After year 3")))', BB, None, TOT_FILL)
    cell(ws, "A16", "Cost of doing nothing, per year", BB)
    cell(ws, "B16", "=Benefits!D10", LINK, NUM)
    cell(ws, "C16", "Every year you wait, this is what stays on the table.", DIM)
    ws.conditional_formatting.add("B8:E8", CellIsRule(operator="lessThan", formula=["0"], font=Font(name=F, bold=True, color="C0392B")))
    widths(ws, [30, 15, 15, 15, 15])

    ws = wb.create_sheet("Rollout plan")
    title(ws, "Rollout plan: how the risk is managed", "The section boards read most closely. The decision covers the whole programme. Each phase has a scope, a measure and a milestone, so progress is visible all the way.")
    header(ws, 4, ["Phase", "Scope", "Duration", "Success measure", "Milestone", "Owner"])
    ph = [("1. Production and quality", "First sites in the rollout plan", "6 to 9 months", "OEE up 3 points against baseline; scrap down 10%", "Live on every shift at the first sites", "Operations director"),
          ("2. Remaining sites", "Production and quality at every other site, from the template", "12 to 24 months", "Go-live on time with under 10% local change", "Each site live on every shift", "Programme lead"),
          ("3. Inventory and maintenance", "Switched on across every site", "6 to 12 months", "Benefits on Benefits sheet reached at 80% ramp", "Live at every site", "Programme lead")]
    for i, row in enumerate(ph, start=5):
        for j, v in enumerate(row):
            inp(ws.cell(row=i, column=1 + j), v)
            ws.cell(row=i, column=1 + j).alignment = WRAP
        ws.row_dimensions[i].height = 32
    for i in range(8, 11):
        for j in range(6):
            inp(ws.cell(row=i, column=1 + j), None)
    widths(ws, [14, 30, 15, 36, 30, 20])
    for s in wb.worksheets:
        s.sheet_properties.tabColor = VIOLET
    return wb


# ---------------------------------------------------------------- swimlane
def swimlane():
    wb = Workbook()
    start_sheet(wb, "Start here", "Swimlane process map", [
        "Map one real order through your factory, from release to the confirmation going back to ERP.",
        "How to use it:",
        "Each row is a lane: a role or a system. Each column is a step in time, left to right. Write a short action in the cell where it happens.",
        "Map what actually happens, not the procedure. Stand on the line and follow the order. Ask the people doing the work.",
        "Mark every hand-off between lanes. Hand-offs are where time gets lost and errors creep in.",
        "Log each pain point on the Pain points sheet as you go, with the time it costs. That list becomes your requirements.",
        "The Example sheet shows a typical as-is. Map yours on Blank map, then design the to-be on a copy.",
    ])
    lanes = ["Planner", "Shift supervisor", "Operator", "Quality", "Warehouse", "Maintenance", "ERP", "Paper / spreadsheet"]
    ex = {
        (0, 1): "Releases order in ERP, prints the pack", (6, 1): "Order released",
        (7, 2): "Traveller, BOM and drawings printed", (1, 2): "Collects pack, walks it to the line",
        (4, 3): "Picks material from pick list", (2, 3): "Waits for material",
        (2, 4): "Sets up machine from paper instruction", (3, 4): "First-off check, signs paper form",
        (2, 5): "Runs order, tallies counts on traveller", (5, 5): "Called by radio when machine stops",
        (2, 6): "Writes scrap and downtime on sheet", (7, 6): "Downtime sheet on clipboard",
        (3, 7): "In-process checks on paper", (1, 7): "Chases quality sign-off",
        (1, 8): "Collects sheets at end of shift", (7, 8): "Supervisor types counts into Excel",
        (0, 9): "Re-keys confirmation into ERP next morning", (6, 9): "Order confirmed, a day late",
    }
    for name, data in (("Example", ex), ("Blank map", {})):
        ws = wb.create_sheet(name)
        title(ws, "Example: a typical as-is, paper-based order" if data else "Your map: order ______ on line ______ , date ______",
              "Lanes down the side, time across the top. One short action per cell. Shade hand-offs." if data else "Copy this sheet to map the to-be once the as-is is agreed.", 10)
        header(ws, 4, ["Lane / step →"] + [f"Step {i}" for i in range(1, 11)])
        for r, lane in enumerate(lanes):
            c = ws.cell(row=5 + r, column=1, value=lane)
            c.font, c.fill, c.border = BB, SUB_FILL, BOX
            ws.row_dimensions[5 + r].height = 46
            for s in range(1, 11):
                v = data.get((r, s))
                c = ws.cell(row=5 + r, column=1 + s, value=v)
                c.font, c.border, c.alignment = (B, BOX, Alignment(wrap_text=True, vertical="center", horizontal="center"))
                if v:
                    c.fill = PatternFill("solid", fgColor="FDE2F3" if lane == "Paper / spreadsheet" else "DDF5F7")
        widths(ws, [20] + [17] * 10)
        if data:
            ws.cell(row=14, column=1, value="Count the hand-offs: in this example the order changes hands 11 times and is re-keyed twice before ERP knows it's done.").font = DIM
    ws = wb.create_sheet("Pain points")
    title(ws, "Pain points log", "One row per problem you see while mapping. Time lost × frequency gives you a number for the business case.", 9)
    header(ws, 4, ["#", "Step", "Lane", "What goes wrong", "Minutes lost each time", "Times per week", "Hours per year", "Idea to fix it", "Priority"])
    dv = DataValidation(type="list", formula1='"High,Medium,Low"', allow_blank=True)
    ws.add_data_validation(dv)
    sample = [1, "Step 9", "Planner", "Confirmations re-keyed from paper the next morning", 20, 15, None, "Confirm automatically from the line", "High"]
    for i in range(5, 31):
        ws.cell(row=i, column=1, value=i - 4).font = B
        for j in range(2, 10):
            c = ws.cell(row=i, column=j)
            if j == 7:
                c.value = f'=IF(AND(ISNUMBER(E{i}),ISNUMBER(F{i})),E{i}*F{i}*48/60,"")'
                c.font, c.number_format, c.border = B, NUM, BOX
            else:
                inp(c, sample[j - 1] if i == 5 else None)
                c.alignment = WRAP
        dv.add(f"I{i}")
    cell(ws, "F32", "Total", BB)
    cell(ws, "G32", "=SUM(G5:G30)", BB, NUM, TOT_FILL)
    ws.cell(row=33, column=1, value="Hours per year assumes 48 working weeks. Row 1 is an example. Overwrite it.").font = DIM
    widths(ws, [5, 10, 16, 44, 12, 11, 12, 36, 10])
    return wb


# ---------------------------------------------------------------- first line
def first_line():
    wb = Workbook()
    start_sheet(wb, "Start here", "Choosing the first line: a scorecard", [
        "Score each candidate line from 1 to 5 on each question. The highest total is usually the right place to start, but read the cautions before you decide.",
        "How to use it:",
        "Put your candidate lines across the top of the Scorecard sheet. Agree the scores as a group: operations, quality, maintenance and IT in the same room.",
        "Use evidence where you have it. A downtime log beats a feeling.",
        "Two cautions:",
        "Don't pick a line purely because it's the worst performer if its problems are mostly mechanical. A system won't fix a worn gearbox.",
        "Don't pick the easiest line if nothing you learn there applies anywhere else. The first line is also your template.",
    ])
    ws = wb.create_sheet("Scorecard")
    title(ws, "First-line scorecard", "1 = poor fit, 5 = strong fit. Weights default to 1. Raise one to 2 only if the whole team agrees it matters more at your site.", 11)
    qs = [("How much does its downtime or scrap cost?", "Value you can measure is what builds momentum for the programme."),
          ("Is it the bottleneck, or close to it?", "An hour gained on the bottleneck is an hour gained for the whole site."),
          ("Is the process reasonably stable and documented?", "You want to prove the system, not fight the process at the same time."),
          ("Is the supervisor keen?", "A champion on the line makes adoption much faster."),
          ("Can the equipment be connected?", "Automatic data capture is where the biggest early wins come from."),
          ("Is it representative of other lines?", "What you learn should carry over when you scale."),
          ("Is the master data in decent shape, or quick to fix?", "Poor master data is one of the most common causes of delay.")]
    lines = ["Line 1", "Line 2", "Line 3", "Line 4", "Line 5", "Line 6", "Line 7", "Line 8"]
    header(ws, 4, ["Question", "Why it matters", "Weight"] + lines)
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True)
    dv.error, dv.errorTitle = "Enter a whole number from 1 to 5.", "Score 1 to 5"
    ws.add_data_validation(dv)
    for c in range(4, 12):
        inp(ws.cell(row=4, column=c), lines[c - 4])
        ws.cell(row=4, column=c).font = Font(name=F, size=10, bold=True, color="0000FF")
    example = [[4, 3, 2, 3], [5, 2, 3, 2], [3, 4, 2, 4], [5, 3, 2, 4], [4, 4, 2, 3], [4, 3, 4, 2], [3, 4, 3, 3]]
    for i, (q, why) in enumerate(qs, start=5):
        cell(ws, f"A{i}", q).alignment = WRAP
        cell(ws, f"B{i}", why, DIM).alignment = WRAP
        inp(ws[f"C{i}"], 1)
        for c in range(4, 12):
            inp(ws.cell(row=i, column=c), example[i - 5][c - 4] if c - 4 < 4 else None)
            dv.add(ws.cell(row=i, column=c))
        ws.row_dimensions[i].height = 32
    cell(ws, "A12", "Weighted total", BB)
    cell(ws, "A13", "Score out of 100", BB)
    cell(ws, "A14", "Rank", BB)
    for c in range(4, 12):
        col = get_column_letter(c)
        cell(ws, f"{col}12", f'=IF(COUNT({col}5:{col}11)=0,"",SUMPRODUCT($C$5:$C$11,{col}5:{col}11))', BB, "0", TOT_FILL)
        cell(ws, f"{col}13", f'=IF({col}12="","",ROUND({col}12/(5*SUM($C$5:$C$11))*100,0))', BB, "0", TOT_FILL)
        cell(ws, f"{col}14", f'=IF({col}12="","",RANK({col}12,$D$12:$K$12))', BB, "0", TOT_FILL)
    ws.conditional_formatting.add("D14:K14", CellIsRule(operator="equal", formula=["1"], fill=PatternFill("solid", fgColor="B9F0C9"), font=Font(name=F, bold=True, color="0B6B2E")))
    cell(ws, "A16", "Notes and evidence", BB)
    for r in range(17, 22):
        for c in range(1, 12):
            inp(ws.cell(row=r, column=c), None)
    ws.cell(row=23, column=1, value="Lines 1 to 4 hold example scores. Overwrite them with your own.").font = DIM
    widths(ws, [40, 46, 8] + [10] * 8)
    return wb


# ---------------------------------------------------------------- go-live
def go_live():
    wb = Workbook()
    start_sheet(wb, "Start here", "Go-live and hypercare checklist", [
        "Go-live is the starting line, not the finish. What happens in the first few weeks decides whether people trust the system.",
        "How to use it:",
        "At the start of the final phase, give every item on Checklist an owner and a due date. Review it weekly with the project team, then daily in the last fortnight.",
        "If something doesn't have an owner, it doesn't get done.",
        "Agree who makes the go / no-go call, and on what evidence, before go-live week. Not at midnight on cutover day.",
        "Use First fortnight to plan who is on the floor each shift, and Issue log from the first hour of production.",
    ])
    ws = wb.create_sheet("Checklist")
    title(ws, "Readiness checklist", "Status: Not started, In progress, Done, or N/A. The summary at the top updates as you go.", 7)
    items = {
        "Scope and sign-off": ["Go-live scope frozen and signed off: lines, processes, users and shifts", "To-be processes signed off by their owners", "Open change requests decided: in, deferred or rejected, with a reason", "Go / no-go decision-maker and criteria agreed in writing"],
        "Master data": ["Materials, BOMs, routings and work centres checked against the shopfloor", "Routings use real cycle times, not planning estimates", "Opening stock and WIP counted and loaded", "Users, roles and skills loaded and access tested for each role", "Quality plans, specs and limits loaded and checked by quality", "Downtime and scrap reason codes agreed. Short lists people will use", "Every label template printed on the real printers and scanned"],
        "Integration": ["ERP orders down and confirmations up, tested end to end with realistic volumes", "Exceptions tested: changed orders, quantity differences, duplicate messages", "Interface monitoring live, with a named owner for failed messages", "Machine connectivity validated machine by machine against what the machine did", "Recovery after a network drop tested. Nothing lost or duplicated"],
        "Testing": ["User acceptance signed off by operational users, including edge cases", "Performance tested at the busiest shift, not a quiet afternoon", "Site acceptance on the real shopfloor: real devices, real network, real gloves", "Cutover rehearsed end to end, with timings", "No severity 1 or 2 defects open. Workarounds written for the rest"],
        "Devices and infrastructure": ["Production environment sized, secured, backed up and monitored", "Every device installed and tested at the point of use", "Wi-Fi coverage tested everywhere it's needed, including the far corners", "Backup and restore actually tested", "Spare devices on hand"],
        "People and training": ["Super users trained deeply and released from their day job for hypercare", "Every operator on every shift trained hands-on, nights and weekends included", "Training and quick guides in the languages people speak", "Supervisors trained on dashboards, approvals and deviations", "Administrators able to make small configuration changes", "What changes explained role by role, not as 'digital transformation'", "Known concerns heard and answered, with the sponsor visibly backing it"],
        "Hypercare and support": ["Hypercare plan: how long, who's on the floor each shift, where the war room is", "Vendor or partner support booked, with people who know your configuration", "One issue log, agreed severity definitions and a daily triage time", "Escalation path published", "Hypercare exit criteria agreed, so it ends on evidence, not a date"],
        "Cutover and rollback": ["Cutover plan written hour by hour, with owners", "Cutover timed for a quieter period, not month end or peak season", "Decision made on in-flight orders: finish the old way or migrate", "Production plan eased for the first days", "Rollback criteria, deadline and procedure agreed and rehearsed", "Paper forms or the old system ready for the fallback period"],
        "KPIs and communication": ["Baseline KPIs captured before go-live, using fixed definitions", "Morning meeting dashboard ready for day one", "Every shift briefed in person on what changes and who to call", "Leadership told what week one looks like: slower and messier at first", "Benefits tracking owner and review dates set"],
    }
    total = sum(len(v) for v in items.values())
    cell(ws, "A4", "Items", BB)
    cell(ws, "B4", total, BB)
    labels = [("Done", "Done"), ("In progress", "In progress"), ("Not started", "Not started"), ("N/A", "N/A")]
    for j, (lab, v) in enumerate(labels):
        cell(ws, f"{get_column_letter(3+2*j)}4", lab, BB)
        cell(ws, f"{get_column_letter(4+2*j)}4", f'=COUNTIF($E$8:$E${7+total+len(items)},"{v}")', BB, "0", TOT_FILL)
    cell(ws, "A5", "Ready (done or N/A)", BB)
    cell(ws, "B5", f'=IF(B4=0,0,(D4+J4)/B4)', BB, PCT, TOT_FILL)
    header(ws, 7, ["Area", "Check", "Owner", "Due", "Status", "Evidence / notes", "Must-have?"])
    dv = DataValidation(type="list", formula1='"Not started,In progress,Done,N/A"', allow_blank=True)
    dv2 = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    ws.add_data_validation(dv)
    ws.add_data_validation(dv2)
    r = 8
    for area, lst in items.items():
        c = ws.cell(row=r, column=1, value=area)
        c.font, c.fill = BB, SUB_FILL
        for j in range(2, 8):
            ws.cell(row=r, column=j).fill = SUB_FILL
        r += 1
        for it in lst:
            cell(ws, f"A{r}", area, DIM)
            cell(ws, f"B{r}", it).alignment = WRAP
            for col in "CDF":
                inp(ws[f"{col}{r}"], None)
            ws[f"D{r}"].number_format = "dd mmm yyyy"
            inp(ws[f"E{r}"], "Not started")
            inp(ws[f"G{r}"], "Yes")
            dv.add(f"E{r}")
            dv2.add(f"G{r}")
            ws.row_dimensions[r].height = 28
            r += 1
    ws.conditional_formatting.add(f"E8:E{r}", CellIsRule(operator="equal", formula=['"Done"'], fill=PatternFill("solid", fgColor="B9F0C9")))
    ws.conditional_formatting.add(f"E8:E{r}", CellIsRule(operator="equal", formula=['"In progress"'], fill=PatternFill("solid", fgColor="FFE8B5")))
    ws.freeze_panes = "C8"
    widths(ws, [22, 62, 16, 13, 14, 36, 11])

    ws = wb.create_sheet("Go no-go")
    title(ws, "Go / no-go decision", "Run this at the go / no-go meeting with the sponsor in the room. Every line needs evidence, not a feeling.", 4)
    header(ws, 4, ["Question", "Yes / No", "Evidence", "Who confirmed"])
    gq = ["Any must-have items on Checklist not done?", "Any severity 1 or 2 defects open?", "Acceptance signed off by operational users?", "Cutover rehearsal completed inside the window?",
          "Master data and opening stock spot-checked in the last 48 hours?", "Every go-live shift trained, with super users rostered?", "Hypercare and vendor support confirmed, with names?",
          "Rollback ready, with a decision point agreed?", "Baseline KPIs captured?"]
    dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    ws.add_data_validation(dv)
    for i, q in enumerate(gq, start=5):
        cell(ws, f"A{i}", q).alignment = WRAP
        inp(ws[f"B{i}"], None)
        dv.add(f"B{i}")
        inp(ws[f"C{i}"], None)
        inp(ws[f"D{i}"], None)
        ws.row_dimensions[i].height = 26
    cell(ws, "A15", "Must-have items on Checklist not yet done or N/A", BB)
    cell(ws, "B15", f'=COUNTIFS(Checklist!$G$8:$G${r},"Yes",Checklist!$E$8:$E${r},"<>Done",Checklist!$E$8:$E${r},"<>N/A")', LINK, "0", TOT_FILL)
    cell(ws, "A16", "Suggested call", BB)
    cell(ws, "B16", '=IF(OR(B5="Yes",B6="Yes",B15>0),"No-go, or go with written conditions",IF(COUNTIF(B7:B13,"No")>0,"Go with conditions","Go"))', BB, None, TOT_FILL)
    ws.merge_cells("B16:D16")
    cell(ws, "A18", "Decision (Go / Go with conditions / No-go)", BB)
    inp(ws["B18"], None)
    cell(ws, "A19", "Conditions, owners and dates", BB)
    inp(ws["B19"], None)
    ws.merge_cells("B19:D19")
    cell(ws, "A20", "Signed by the sponsor, and date", BB)
    inp(ws["B20"], None)
    ws.cell(row=22, column=1, value="A delayed go-live is uncomfortable for a week. A failed one costs you the operators' trust for months.").font = DIM
    widths(ws, [52, 34, 40, 20])

    ws = wb.create_sheet("First fortnight")
    title(ws, "The first fortnight on the floor", "Who is where, every shift. Leaders on the shopfloor, not in the war room.", 6)
    header(ws, 4, ["When", "Focus", "Day shift", "Back shift", "Night shift", "Done?"])
    plan = [("Cutover", "Run the rehearsed plan. First order released, first label printed, first confirmation in ERP."),
            ("Day 1", "Super users and project team on every shift, visible. Log every issue, however small. Fix quick ones the same day."),
            ("Days 2 and 3", "Daily triage of the issue log. Watch the interfaces under real load, and night shift especially."),
            ("Days 4 and 5", "Close the issues that erode trust first. Tell the floor what's been fixed."),
            ("End of week 1", "Review against exit criteria. Celebrate the first wins. Honest update to the steering group and the floor."),
            ("Week 2", "Fewer people on the floor if the log allows. Check the first reports against reality."),
            ("End of week 2", "Exit hypercare on evidence, or extend it on purpose. Start comparing KPIs with the baseline.")]
    for i, (w, f) in enumerate(plan, start=5):
        cell(ws, f"A{i}", w, BB)
        cell(ws, f"B{i}", f).alignment = WRAP
        for col in "CDEF":
            inp(ws[f"{col}{i}"], None)
        ws.row_dimensions[i].height = 40
    widths(ws, [16, 60, 18, 18, 18, 9])

    ws = wb.create_sheet("Issue log")
    title(ws, "Hypercare issue log", "One place for every issue. Severity 1 stops production, 2 has a painful workaround, 3 is an annoyance, 4 is a request.", 9)
    header(ws, 4, ["#", "Raised", "Line / area", "Shift", "What happened", "Severity", "Owner", "Status", "Fix and date"])
    dv = DataValidation(type="list", formula1='"1,2,3,4"', allow_blank=True)
    dv2 = DataValidation(type="list", formula1='"Open,Fixing,Fixed,Closed"', allow_blank=True)
    ws.add_data_validation(dv)
    ws.add_data_validation(dv2)
    for i in range(5, 105):
        ws.cell(row=i, column=1, value=i - 4).font = B
        for j in range(2, 10):
            inp(ws.cell(row=i, column=j), None)
        ws.cell(row=i, column=2).number_format = "dd mmm hh:mm"
        dv.add(f"F{i}")
        dv2.add(f"H{i}")
    ws.conditional_formatting.add("A5:I104", FormulaRule(formula=['AND($F5=1,$H5<>"Closed",$H5<>"Fixed")'], fill=PatternFill("solid", fgColor="FAD4D4")))
    ws.freeze_panes = "A5"
    widths(ws, [5, 14, 16, 10, 50, 9, 14, 10, 36])
    return wb


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    for fn, name in [(business_case, "shopfloor-matters-business-case-template.xlsx"), (swimlane, "shopfloor-matters-swimlane-process-map.xlsx"),
                     (first_line, "shopfloor-matters-first-line-scorecard.xlsx"), (go_live, "shopfloor-matters-go-live-hypercare-checklist.xlsx")]:
        wb = fn()
        for ws in wb.worksheets:
            ws.sheet_properties.tabColor = VIOLET
            ws.page_setup.orientation = "landscape"
            ws.page_setup.fitToWidth = 1
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.page_setup.fitToHeight = 0
            ws.oddFooter.left.text = "Shopfloor Matters · shopfloormatters.com"
        wb.properties.creator = "Matt Barber, Shopfloor Matters"
        wb.save(OUT / name)
        print("wrote", name)
